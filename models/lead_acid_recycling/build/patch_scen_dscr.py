# -*- coding: utf-8 -*-
"""یکسان‌سازیِ مبنایِ DSCR در مدلِ سایهٔ سناریو (ENG_SCEN) با مدلِ اصلی (ENG_FCF).

علتِ اصلاح: دو سطرِ CFADS (۶۵/۹۹/۱۳۳) و خدمت بدهی (۶۶/۱۰۰/۱۳۴) در مدلِ سایه بر
مبنایِ متفاوتی از مدلِ اصلی نوشته شده بودند:

  · CFADSِ سایه از مالیاتِ «روی EBIT» استفاده می‌کرد (بدون سپرِ مالیاتیِ بهره)،
    در حالی که مدلِ اصلی مالیات را روی EBT می‌بندد:
        EBT = EBIT − سود بلندمدت − کارمزد − سود کوتاه‌مدت + درآمد بهره
  · خدمت بدهیِ سایه سود را به‌صورت «ماندهٔ اول دوره × نرخ» می‌گرفت و سودِ
    تسهیلاتِ کوتاه‌مدتِ سرمایه در گردش را حذف می‌کرد، در حالی که مدلِ اصلی
    خدمت بدهی را «اصل + سود بلندمدت + کارمزد + سود کوتاه‌مدت» تعریف می‌کند.

اثرِ این دو تفاوت: سناریویِ پایهٔ مدلِ سایه با مدلِ اصلی یکی نبود (مثلاً کمینهٔ
DSCR برابر ‎−۱.۳۶‎ در برابر ‎−۱.۲۵‎) و DSCRِ سناریوها در سال‌های پایانی تا چندصد
برابر بزرگ‌نمایی می‌شد. پس از اصلاح، سناریوی پایه در همهٔ دوره‌ها دقیقاً برابرِ
مدلِ اصلی است و سناریوهای دیگر هم بر همان مبنا محاسبه می‌شوند.

نکته: سطرهای مربوط به FCFF/NPV/IRR دست‌نخورده می‌مانند (مبنایِ آن‌ها ناپوشیده و
درست است)، بنابراین نتایجِ ارزش‌گذاری تغییر نمی‌کند؛ تنها سطرهای CFADS، خدمت
بدهی و در نتیجه DSCR و کمینه/میانگینِ آن به‌روز می‌شوند.

اجرا:  python build/patch_scen_dscr.py [--apply]
"""

import os
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from mini_excel import Evaluator, XLError  # noqa: E402
import spec  # noqa: E402

BLOCK_ROWS = {"cfads": 65, "debt_service": 66}
BLOCK_STEP = 34          # فاصلهٔ بلوک‌های سناریو
SCEN_COL = {0: "D", 1: "E", 2: "F"}


def kratio(scen):
    """نسبتِ نرخ بهرهٔ سناریو به نرخِ سناریوی فعال"""
    return "('IN_FIN'!$D$25+'IN_SCEN'!$%s$21)/('IN_FIN'!$D$25+'ENG_SCEN'!$G$19)" % SCEN_COL[scen]


def new_formulas(scen, L):
    o = BLOCK_STEP * scen
    kr = kratio(scen)
    debt = ("='ENG_FCF'!%s$24+'ENG_FCF'!%s$26*%s+'ENG_FCF'!%s$27+'ENG_FCF'!%s$28"
            % (L, L, kr, L, L))
    cfads = ("='ENG_SCEN'!%s%d-MAX(0,'ENG_SCEN'!%s%d-('ENG_FCF'!%s$26*%s+'ENG_FCF'!%s$27"
             "+'ENG_FCF'!%s$28-'ENG_FCF'!%s$29))*'IN_MACRO'!$D$14-'ENG_SCEN'!%s%d"
             % (L, 49 + o, L, 50 + o, L, kr, L, L, L, L, 54 + o))
    return cfads, debt


def main(apply=False):
    path = spec.MODEL_PATH
    wb = load_workbook(path)
    ws = wb["ENG_SCEN"]
    print("مبنای پیشین:")
    for i, base in ((0, "پایه"), (1, "خوش‌بینانه"), (2, "بدبینانه")):
        o = BLOCK_STEP * i
        print("  %-10s CFADS ردیف %-3d %s" % (base, 65 + o, ws.cell(65 + o, 4).value))
        print("  %-10s خدمت   ردیف %-3d %s" % ("", 66 + o, ws.cell(66 + o, 4).value))

    n = 0
    if apply:
        for i in range(3):
            o = BLOCK_STEP * i
            for t in range(spec.NY):
                L = get_column_letter(spec.C0 + t)
                cf, ds = new_formulas(i, L)
                ws.cell(BLOCK_ROWS["cfads"] + o, spec.C0 + t).value = cf
                ws.cell(BLOCK_ROWS["debt_service"] + o, spec.C0 + t).value = ds
                n += 2
        wb.save(path)
        print("\n✓ %d سلول اصلاح شد" % n)

    # ---- بررسی: سناریوی پایهٔ سایه باید برابرِ مدل اصلی شود
    ev = Evaluator(spec.MODEL_PATH)
    print("\nبررسیِ انطباقِ DSCR سناریوی پایه با مدل اصلی (باید برابر باشد):")
    bad = 0
    for t in range(spec.NY):
        a = ev.get("ENG_FCF", 44, spec.C0 + t)
        b = ev.get("ENG_SCEN", 67, spec.C0 + t)
        if isinstance(a, XLError) or isinstance(b, XLError):
            print("   t=%-2d خطا" % t)
            bad += 1
            continue
        ok = abs(a - b) < 1e-9
        bad += 0 if ok else 1
        print("   t=%-2d اصلی %10.4f | سایه %10.4f  %s" % (t, a, b, "✓" if ok else "✗"))
    print("→ %s" % ("منطبق ✅" if not bad else "%d دوره ناهماهنگ ❌" % bad))
    return 0 if not bad else 1


if __name__ == "__main__":
    sys.exit(main(apply="--apply" in sys.argv))
