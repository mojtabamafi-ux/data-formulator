# -*- coding: utf-8 -*-
"""زیرساختِ مشترکِ ساختِ کاربرگ‌ها.

دو ایدهٔ اصلی اینجا پیاده شده است:

۱) آدرس‌دهی با کلید، نه با شمارهٔ سطر. هر سطرِ داده با یک کلید (ستونِ B) شناخته
   می‌شود؛ بنابراین اگر ترتیبِ سطرها عوض شود، فرمول‌ها همچنان درست می‌مانند و
   ابزارهای گزارش‌گیری (spec.discover_keys) بدون تغییر کار می‌کنند.

۲) تعویقِ ارجاع‌ها. هنگام نوشتنِ فرمول ممکن است شیت یا کلیدِ مقصد هنوز ساخته
   نشده باشد (مثلاً ENG_SCEN به ENG_OPS ارجاع می‌دهد که بعداً ساخته می‌شود).
   در این حالت ref() یک نشانِ موقت (@@SH|KEY|t@@) برمی‌گرداند و در پایان،
   resolve() همهٔ آن‌ها را به آدرسِ واقعی تبدیل می‌کند.

قاعدهٔ مدل: شیت‌های ورودی فقط عدد و متن دارند (بدون فرمول)؛ هر سلولِ لایهٔ موتور
و خروجی باید فرمول باشد.
"""

import re

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

import spec

# ------------------------------------------------------------------ ظاهر
TITLE_FONT = Font(bold=True, size=14, color="1F3864")
SUB_FONT = Font(size=9, italic=True, color="595959")
SECTION_FONT = Font(bold=True, size=11, color="1F3864")
SECTION_FILL = PatternFill("solid", fgColor="D9E1F2")
HEAD_FILL = PatternFill("solid", fgColor="F2F2F2")
KEY_FONT = Font(size=9, name="Consolas", color="404040")
UNIT_FONT = Font(size=9, color="808080")
NOTE_FONT = Font(size=9, italic=True, color="7F7F7F")

TAB_COLOR = {"input": "2E75B6", "engine": "A6A6A6", "output": "548235",
             "check": "C00000", "map": "1F3864"}
LAYER_LABEL = {"input": "لایه ورودی", "engine": "موتور محاسبات",
               "output": "لایه خروجی", "check": "کنترل", "map": "نقشه و راهنما"}

THIN = Side(style="thin", color="BFBFBF")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
INPUT_FILL = PatternFill("solid", fgColor="FFF2CC")   # سلول‌های ورودی: زرد کمرنگ

FMT = {
    "num":   "#,##0",
    "num1":  "#,##0.0",
    "num2":  "#,##0.00",
    "num3":  "#,##0.000",
    "pct":   "0.0%",
    "pct2":  "0.00%",
    "idx":   "0",
    "mult":  "0.0",
    "text":  "General",
}

PLACEHOLDER = re.compile(r"@@([^|]+)\|([^|]*)\|(-?\d+)@@")


class Sheet:
    """یک کاربرگ با مکان‌نمای سطر و ابزارهای نوشتنِ ردیف‌های کلیددار"""

    def __init__(self, book, name, title, subtitle, layer, time_axis=True,
                 width_a=54):
        self.book = book
        self.name = name
        self.layer = layer
        self.ws = book.wb.create_sheet(name)
        self.r = 1
        self.ws.sheet_view.rightToLeft = True
        self.ws.sheet_properties.tabColor = TAB_COLOR.get(layer, "A6A6A6")
        self.ws.column_dimensions["A"].width = width_a
        self.ws.column_dimensions["B"].width = 14
        self.ws.column_dimensions["C"].width = 18
        for i in range(spec.NY):
            self.ws.column_dimensions[get_column_letter(spec.C0 + i)].width = 14
        self.ws.freeze_panes = "D7"

        self.ws.cell(1, 1, title).font = TITLE_FONT
        c = self.ws.cell(1, 2, LAYER_LABEL.get(layer, ""))
        c.font = Font(bold=True, size=9, color="595959")
        self.ws.cell(2, 1, subtitle).font = SUB_FONT
        self.r = 4
        if time_axis:
            self.r = 7
            self._axis()

    def _axis(self):
        """سطرهای ۷ و ۸: محور زمان (دوره و سال تقویمی)."""
        ws = self.ws
        c = ws.cell(7, 1, "دوره (t)")
        c.font = Font(bold=True, size=9)
        c.fill = HEAD_FILL
        c2 = ws.cell(8, 1, "سال تقویمی")
        c2.font = Font(bold=True, size=9)
        c2.fill = HEAD_FILL
        for t in range(spec.NY):
            col = spec.C0 + t
            if self.layer == "input":                # ورودی‌ها بی‌فرمول‌اند
                ws.cell(7, col, t).number_format = FMT["idx"]
                ws.cell(8, col, self.book.start_year + t).number_format = FMT["idx"]
            else:                                    # موتور/خروجی: ارجاع به IN_MACRO
                ws.cell(7, col, self.book.ref("IN_MACRO", "MAC.t", t)).number_format = FMT["idx"]
                ws.cell(8, col, self.book.ref("IN_MACRO", "MAC.year", t)).number_format = FMT["idx"]
            ws.cell(7, col).fill = HEAD_FILL
            ws.cell(8, col).fill = HEAD_FILL
        self.r = 9

    # ------------------------------------------------------------ ابزارها
    def gap(self, n=1):
        self.r += n
        return self.r

    def section(self, text, row=None):
        r = row or self.r
        c = self.ws.cell(r, 1, text)
        c.font = SECTION_FONT
        c.fill = SECTION_FILL
        for col in range(2, spec.C0 + spec.NY):
            self.ws.cell(r, col).fill = SECTION_FILL
        if row is None:
            self.r = r + 1
        return r

    def note(self, text, row=None):
        r = row or self.r
        c = self.ws.cell(r, 1, text)
        c.font = NOTE_FONT
        c.alignment = Alignment(wrap_text=False)
        if row is None:
            self.r = r + 1
        return r

    def header(self, items, row=None, col0=4):
        """یک سطرِ سرستون برای جدول‌های ایستگاهی/محصولی"""
        r = row or self.r
        for i, txt in enumerate(items):
            c = self.ws.cell(r, col0 + i, txt)
            c.font = Font(bold=True, size=9)
            c.fill = HEAD_FILL
            c.border = BOX
        if row is None:
            self.r = r + 1
        return r

    def _label(self, r, key, label, unit):
        self.ws.cell(r, 1, label)
        if key:
            c = self.ws.cell(r, 2, key)
            c.font = KEY_FONT
            self.book.keys[(self.name, key)] = r
        if unit:
            c = self.ws.cell(r, 3, unit)
            c.font = UNIT_FONT
            c.alignment = Alignment(horizontal="center")

    def add(self, key, label, unit, values, fmt="num", row=None):
        """یک ردیفِ داده به طولِ افق مدل (ستون‌های D تا P)"""
        r = row or self.r
        self._label(r, key, label, unit)
        for t, v in enumerate(values):
            if v is None:
                continue
            c = self.ws.cell(r, spec.C0 + t, v)
            c.number_format = FMT[fmt]
            if self.layer == 0:
                c.fill = INPUT_FILL
                c.border = BOX
        if row is None:
            self.r = r + 1
        return r

    def add_scalar(self, key, label, unit, value, fmt="num", row=None):
        """یک ردیف با مقدارِ واحد در ستونِ D"""
        return self.add(key, label, unit, [value], fmt=fmt, row=row)

    def add_row(self, key, label, unit, fmt="num", row=None):
        """فقط تخصیصِ سطر (برای پر شدن با فرمول در مرحلهٔ بعد)"""
        r = row or self.r
        self._label(r, key, label, unit)
        for t in range(spec.NY):
            self.ws.cell(r, spec.C0 + t).number_format = FMT[fmt]
        if row is None:
            self.r = r + 1
        return r

    def local_axis(self, row=None):
        """محور زمانِ محلی برای جدول‌های میانِ شیت (مثل منحنی دستیابی به ظرفیت)"""
        r = row or self.r
        for i, lab in ((0, "دوره (t)"), (1, "سال تقویمی")):
            c = self.ws.cell(r + i, 1, lab)
            c.font = Font(bold=True, size=9)
            c.fill = HEAD_FILL
            for t in range(spec.NY):
                col = spec.C0 + t
                if self.layer == "input":
                    v = t if i == 0 else self.book.start_year + t
                    self.ws.cell(r + i, col, v).number_format = FMT["idx"]
                else:
                    key = "MAC.t" if i == 0 else "MAC.year"
                    self.ws.cell(r + i, col,
                                 self.book.ref("IN_MACRO", key, t)).number_format = FMT["idx"]
                self.ws.cell(r + i, col).fill = HEAD_FILL
        if row is None:
            self.r = r + 2
        return r

    def table(self, key, label, unit, values, fmt="num", row=None, n=None, col0=4):
        """یک ردیفِ جدولی (مثلِ پارامترهای ایستگاه‌ها) با کلید در ستونِ B"""
        r = row or self.r
        self._label(r, key, label, unit)
        for i, v in enumerate(values[: (n or len(values))]):
            c = self.ws.cell(r, col0 + i, v)
            c.number_format = FMT[fmt]
            if self.layer == 0:
                c.fill = INPUT_FILL
                c.border = BOX
        if row is None:
            self.r = r + 1
        return r


class Book:
    """کتابِ کار: نگه‌داریِ کلیدها، نشان‌های موقت و نوشتنِ فرمول‌ها"""

    def __init__(self, start_year=2026):
        self.wb = Workbook()
        self.wb.remove(self.wb.active)
        self.keys = {}          # {(شیت، کلید): سطر}
        self.sheets = {}        # {نام شیت: Sheet}
        self.start_year = start_year
        self._pending = set()

    # ------------------------------------------------------- ساختِ کاربرگ
    def sheet(self, name, title, subtitle, layer, **kw):
        sh = Sheet(self, name, title, subtitle, layer, **kw)
        self.sheets[name] = sh
        return sh

    # ---------------------------------------------------------- آدرس‌دهی
    def row_of(self, sheet, key):
        return self.keys[(sheet, key)]

    def ref(self, sheet, key, t=0, abs_row=True, abs_col=True):
        """آدرسِ مطلقِ یک سلول بر پایهٔ کلید و شمارهٔ دوره"""
        try:
            r = self.keys[(sheet, key)]
        except KeyError:
            self._pending.add((sheet, key))
            return "@@%s|%s|%d@@" % (sheet, key, t)
        col = get_column_letter(spec.C0 + (t or 0))
        return "'%s'!%s%s%s%d" % (sheet, "$" if abs_col else "", col,
                                  "$" if abs_row else "", r)

    def rng(self, sheet, key, t0=0, t1=None):
        """آدرسِ یک بازهٔ دوره‌ای از یک ردیف"""
        t1 = spec.NY - 1 if t1 is None else t1
        try:
            r = self.keys[(sheet, key)]
        except KeyError:
            self._pending.add((sheet, key))
            return "@@%s|%s|%d@@:@@%s|%s|%d@@" % (sheet, key, t0, sheet, key, t1)
        c0 = get_column_letter(spec.C0 + t0)
        c1 = get_column_letter(spec.C0 + t1)
        return "'%s'!$%s$%d:$%s$%d" % (sheet, c0, r, c1, r)

    def cellref(self, sheet, row, col):
        return "'%s'!$%s$%d" % (sheet, get_column_letter(col), row)

    def src(self, sheet, key, t=0):
        """نامِ شیت به‌صورتِ کوتاه برای نمایش (برای برچسب‌های گزارش)"""
        return "%s.%s[%d]" % (sheet, key, t)

    # ----------------------------------------------------------- فرمول‌ها
    def f(self, sheet, key, t, formula):
        """نوشتنِ فرمول در خانهٔ (کلید، دوره)"""
        r = self.keys[(sheet, key)]
        self.wb[sheet].cell(r, spec.C0 + (t or 0), formula)

    def frow(self, sheet, key, tmpl, rng_=range(spec.NY)):
        """نوشتنِ فرمول برای همهٔ دوره‌ها؛ tmpl یک تابع از t است"""
        r = self.keys[(sheet, key)]
        for t in rng_:
            self.wb[sheet].cell(r, spec.C0 + t, tmpl(t))

    def fscalar(self, sheet, key, formula, col=None):
        r = self.keys[(sheet, key)]
        self.wb[sheet].cell(r, col or spec.C0, formula)

    def fseed(self, sheet, key, formula):
        """نوشتنِ فرمول در نخستین ستونِ داده (مبالغِ واحد)"""
        self.fscalar(sheet, key, formula)

    def ftable(self, sheet, key, tmpl, n, col0=4):
        """نوشتنِ فرمول در یک ردیفِ جدولی (مثلِ طبقاتِ CAPEX)"""
        r = self.keys[(sheet, key)]
        for i in range(n):
            self.wb[sheet].cell(r, col0 + i, tmpl(i))

    def fcell(self, sheet, row, col, formula):
        """نوشتنِ فرمول در یک مختصاتِ صریح (برای سطرهای بدون کلید)"""
        self.wb[sheet].cell(row, col, formula)

    def vcell(self, sheet, row, col, value, fmt="num"):
        c = self.wb[sheet].cell(row, col, value)
        c.number_format = FMT[fmt]
        return c

    # ---------------------------------------------------------- حلِ ارجاع‌ها
    def resolve(self):
        """تبدیلِ نشان‌های موقت به آدرسِ واقعی در همهٔ سلول‌ها"""
        if self._pending:
            missing = [k for k in self._pending if k not in self.keys]
            if missing:
                raise KeyError("کلیدهای تعریف‌نشده: %s" % missing)
        n = 0
        for ws in self.wb.worksheets:
            for row in ws.iter_rows():
                for c in row:
                    v = c.value
                    if isinstance(v, str) and "@@" in v:
                        c.value = PLACEHOLDER.sub(self._resolve_one, v)
                        n += 1
        return n

    def _resolve_one(self, m):
        sheet, key, t = m.group(1), m.group(2), int(m.group(3))
        r = self.keys[(sheet, key)]
        col = get_column_letter(spec.C0 + t)
        return "'%s'!$%s$%d" % (sheet, col, r)

    def save(self, path):
        self.resolve()
        self.wb.save(path)
        return path
