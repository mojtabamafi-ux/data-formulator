# -*- coding: utf-8 -*-
"""مقایسهٔ فایلِ ساخته‌شده با فایلِ مرجع (تحویلی)، سلول‌به‌سلول.

هدف: اطمینان از این‌که مولدِ کد، همان مدل را بازتولید می‌کند. مقایسه بر پایهٔ
محتوای سلول (فرمول یا مقدار) و به‌تفکیکِ شیت انجام می‌شود.

اجرا:  python build/compare_workbook.py [ساخته‌شده] [مرجع]
خروجیِ برنامه: ۰ = انطباق، ۱ = تفاوت
"""

import os
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from openpyxl import load_workbook  # noqa: E402

import spec  # noqa: E402

DEFAULT_BUILT = os.path.join(spec.OUT_DIR, "built.xlsx")


def norm(v):
    if v is None:
        return ""
    if isinstance(v, str):
        return v.replace(" ", "")
    if isinstance(v, float) and v == int(v):
        return int(v)
    return v


def compare(built_path=DEFAULT_BUILT, ref_path=spec.MODEL_PATH, sheets=None, verbose=True):
    wb1 = load_workbook(built_path, data_only=False)
    wb2 = load_workbook(ref_path, data_only=False)
    sheets = sheets or [s for s in wb1.sheetnames if s in wb2.sheetnames]
    total_diff = 0
    if verbose:
        print("=" * 74)
        print("مقایسهٔ فایل ساخته‌شده با مرجع")
        print("=" * 74)
        print("  ساخته‌شده: %s" % os.path.relpath(built_path, spec.ROOT))
        print("  مرجع     : %s" % os.path.relpath(ref_path, spec.ROOT))
        print("  شیت‌های مقایسه‌شده: %d از %d"
              % (len(sheets), len(wb2.sheetnames)))
        missing = [s for s in wb2.sheetnames if s not in wb1.sheetnames]
        if missing:
            print("  هنوز ساخته نشده (%d): %s" % (len(missing), "، ".join(missing)))

    for name in sheets:
        w1, w2 = wb1[name], wb2[name]
        diffs = []
        for r in range(1, max(w1.max_row, w2.max_row) + 1):
            for c in range(1, max(w1.max_column, w2.max_column) + 1):
                a = norm(w1.cell(r, c).value)
                b_ = norm(w2.cell(r, c).value)
                if a != b_:
                    diffs.append((r, c, w1.cell(r, c).value, w2.cell(r, c).value))
        total_diff += len(diffs)
        if verbose:
            flag = "✓" if not diffs else "✗"
            print("  %-12s %s %d تفاوت" % (name, flag, len(diffs)))
            for r, c, a, b_ in diffs[:6]:
                print("       %s%d: ساخته‌شده=%s | مرجع=%s"
                      % (chr(64 + c) if c < 27 else "?", r,
                         ("%r" % a)[:58], ("%r" % b_)[:58]))
            if len(diffs) > 6:
                print("       … و %d مورد دیگر" % (len(diffs) - 6))
    if verbose:
        print("-" * 74)
        print("جمع تفاوت‌ها: %d → %s" % (total_diff, "منطبق ✅" if not total_diff else "ناهماهنگ ❌"))
        print("=" * 74)
    return total_diff


if __name__ == "__main__":
    a = sys.argv[1] if len(sys.argv) > 1 else DEFAULT_BUILT
    b = sys.argv[2] if len(sys.argv) > 2 else spec.MODEL_PATH
    sys.exit(0 if compare(a, b) == 0 else 1)
