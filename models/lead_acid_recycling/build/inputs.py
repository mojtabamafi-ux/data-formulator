# -*- coding: utf-8 -*-
"""لایهٔ ورودی (۷ شیت): فقط عدد و متن، بدونِ هیچ فرمول.

هر تغییر در مدل باید از اینجا آغاز شود؛ لایه‌های موتور و خروجی همگی به این
شیت‌ها ارجاع می‌دهند. مقدارهای زیر همان فرضیه‌های پایهٔ مدل هستند و با فایل
تحویلی یکسان‌اند (توسط compare_workbook.py کنترل می‌شود).
"""

from openpyxl.worksheet.datavalidation import DataValidation

import common
import spec

# ------------------------------------------------------------------ IN_MACRO
START_YEAR = 2026

MACRO_SCALARS = [
    ("MAC.start_year", "سال شروع مدل (میلادی)", "سال", 2026, "idx"),
    ("MAC.horizon", "افق مدل (سال بهره‌برداری)", "سال", 12, "idx"),
    ("MAC.ops_start", "شاخص سال آغاز بهره‌برداری (t)", "شاخص", 1, "idx"),
    ("MAC.constr_end", "شاخص سال پایان دوره ساخت (t)", "شاخص", 1, "idx"),
    ("MAC.tax", "نرخ مالیات بر درآمد", "درصد", 0.20, "pct"),
    ("MAC.erp", "صرف ریسک بازار سهام (ERP)", "درصد", 0.06, "pct"),
    ("MAC.beta", "بتای سهام", "ضریب", 1.15, "num2"),
    ("MAC.crp", "صرف ریسک کشور (CRP)", "درصد", 0.04, "pct"),
    ("MAC.g_term", "نرخ رشد دائمی اسمی", "درصد", 0.18, "pct"),
    ("MAC.tv_method", "روش ارزش پایانی (1=گوردون/2=ضریب خروج)", "1/2", 2, "idx"),
    ("MAC.exit_mult", "ضریب خروج (EV/EBITDA)", "ضریب", 6, "num1"),
]

MACRO_SERIES = [
    ("MAC.t", "شاخص دوره (t)", "شاخص",
     [0, 1, 2, 3, 4, 5, 6, 7, 8, 9, 10, 11, 12], "idx"),
    ("MAC.year", "سال تقویمی", "سال",
     [START_YEAR + t for t in range(spec.NY)], "idx"),
    ("MAC.infl", "نرخ تورم عمومی", "درصد",
     [0.25, 0.25, 0.23, 0.22, 0.20, 0.20, 0.18, 0.18, 0.16, 0.16, 0.15, 0.15, 0.15], "pct"),
    ("MAC.wage", "نرخ رشد دستمزد", "درصد",
     [0.27, 0.26, 0.25, 0.23, 0.22, 0.21, 0.20, 0.19, 0.18, 0.17, 0.17, 0.16, 0.16], "pct"),
    ("MAC.energy", "نرخ رشد قیمت حامل‌های انرژی", "درصد",
     [0.30, 0.28, 0.26, 0.25, 0.23, 0.22, 0.21, 0.20, 0.20, 0.19, 0.18, 0.18, 0.18], "pct"),
    ("MAC.fx", "نرخ ارز (ریال به ازای دلار)", "ریال/دلار",
     [500000, 620000, 760000, 930000, 1120000, 1340000, 1590000,
      1870000, 2180000, 2510000, 2860000, 3230000, 3610000], "num"),
    ("MAC.lme", "قیمت جهانی سرب (LME)", "دلار/تن",
     [2100, 2150, 2200, 2250, 2250, 2300, 2300, 2350, 2350, 2400, 2400, 2450, 2450], "num"),
    ("MAC.rf", "نرخ سود بدون ریسک اسمی", "درصد",
     [0.24, 0.24, 0.23, 0.23, 0.22, 0.22, 0.21, 0.21, 0.20, 0.20, 0.20, 0.19, 0.19], "pct"),
    ("MAC.gdp", "رشد GDP بخش هدف", "درصد",
     [0.03, 0.03, 0.035, 0.035, 0.04, 0.04, 0.04, 0.035, 0.035, 0.03, 0.03, 0.03, 0.03], "pct"),
    ("MAC.ppi", "شاخص قیمت تولیدکننده", "درصد",
     [0.24, 0.24, 0.22, 0.21, 0.20, 0.19, 0.18, 0.18, 0.17, 0.16, 0.16, 0.15, 0.15], "pct"),
    ("MAC.constr_flg", "پرچم دوره ساخت (1=بله)", "1/0",
     [1, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0], "idx"),
    ("MAC.ops_flg", "پرچم دوره بهره‌برداری (1=بله)", "1/0",
     [0, 1, 1, 1, 1, 1, 1, 1, 1, 1, 1, 1, 1], "idx"),
]

MACRO_NOTES = [
    "۱. شاخص‌های قیمتی از سال پایه (t=0) ساخته می‌شوند؛ مقدار شاخص در t=0 برابر یک است.",
    "۲. نرخ ارز به‌صورت «ریال به ازای هر دلار» وارد می‌شود و مسیر آن مبنای تعدیل اقلام ارزی است.",
    "۳. قیمت جهانی سرب (LME) مبنای قیمت‌گذاری محصولاتِ پیوند‌خورده و نیز قیمت خرید باتری ضایعاتی است.",
    "۴. روش ارزش پایانی: ۱=رشد دائمی (گوردون)، ۲=ضریب خروج (EV/EBITDA).",
]

# ------------------------------------------------------------------- IN_TECH
TECH_SCALARS_A = [
    ("TECH.cap_in", "ظرفیت اسمی ورودی باتری", "تن/سال", 40000, "num"),
    ("TECH.days", "روزهای کاری در سال", "روز", 330, "num"),
    ("TECH.shifts", "شیفت در روز", "شیفت", 3, "num"),
    ("TECH.hshift", "ساعت در شیفت", "ساعت", 8, "num"),
    ("TECH.dt_plan", "توقف برنامه‌ریزی‌شده", "ساعت/سال", 400, "num"),
    ("TECH.dt_unplan", "توقف اضطراری", "ساعت/سال", 300, "num"),
]

TECH_SCALARS_B = [
    ("TECH.pb_paste", "عیار سرب در خمیر", "درصد", 0.72, "pct"),
    ("TECH.pb_metal", "عیار سرب در قطعات فلزی", "درصد", 0.95, "pct"),
    ("TECH.desulf", "راندمان سولفات‌زدایی", "درصد", 0.95, "pct"),
    ("TECH.smelt", "راندمان کوره ذوب و احیا", "درصد", 0.97, "pct"),
    ("TECH.refine", "راندمان تصفیه نهایی", "درصد", 0.99, "pct"),
    ("TECH.pp_rec", "راندمان بازیابی پلی‌پروپیلن", "درصد", 0.92, "pct"),
    ("TECH.na2so4", "ضریب تولید سولفات سدیم", "تن/تن", 0.13, "num2"),
    ("TECH.sp_soft", "سهم سرب خالص از سرب تصفیه‌شده", "درصد", 0.60, "pct"),
    ("TECH.sp_alloy", "سهم آلیاژ از سرب تصفیه‌شده", "درصد", 0.40, "pct"),
]

TECH_SCALARS_C = [
    ("TECH.sh_paste", "خمیر سرب (Paste)", "درصد", 0.38, "pct"),
    ("TECH.sh_metal", "قطعات فلزی سربی (Grid / Posts)", "درصد", 0.26, "pct"),
    ("TECH.sh_pp", "پلی‌پروپیلن", "درصد", 0.09, "pct"),
    ("TECH.sh_elec", "الکترولیت (اسید سولفوریک)", "درصد", 0.12, "pct"),
    ("TECH.sh_res", "جداکننده‌ها و سایر پسماند", "درصد", 0.15, "pct"),
]

STATIONS = [
    "پذیرش، تخلیه و جداسازی الکترولیت",
    "خردایش و جدایش هیدرودینامیکی",
    "سولفات‌زدایی خمیر سرب",
    "کوره ذوب و احیا",
    "تصفیه و آلیاژسازی",
    "شمش‌ریزی، بسته‌بندی و انبار",
]

STATION_ROWS = [
    ("TECH.cap_yr", "ظرفیت اسمی (تن در سال)", "تن/سال",
     [42000, 42000, 17000, 27000, 21000, 30000], "num"),
    ("TECH.cap_hr", "ظرفیت اسمی (تن در ساعت)", "تن/ساعت",
     [6.5, 6.5, 2.6, 4.2, 3.3, 4.7], "num1"),
    ("TECH.avail", "نرخ دسترسی (Availability)", "درصد",
     [0.92, 0.90, 0.93, 0.91, 0.94, 0.95], "pct"),
    ("TECH.yield", "راندمان ایستگاه", "درصد",
     [0.98, 0.97, 0.95, 0.97, 0.99, 0.99], "pct2"),
    ("TECH.scrap", "نرخ ضایعات / دورریز", "درصد",
     [0.005, 0.01, 0.01, 0.005, 0.005, 0.01], "pct2"),
    ("TECH.kwh", "انرژی ویژه الکتریکی", "kWh/تن",
     [8, 22, 18, 95, 40, 12], "num"),
    ("TECH.water", "آب ویژه", "m3/تن",
     [0.1, 0.6, 1.2, 0.25, 0.2, 0.35], "num2"),
]

DRIVER_ROWS = [
    ("THROUGHPUT", "جریان مواد (تن/سال در ظرفیت طراحی)", "تن/سال",
     [40000, 40000, 15200, 24840, 19700, 27900], "num"),
    ("MACHINE_HOURS", "ساعت کار ماشین در سال (طراحی)", "ساعت/سال",
     [6000, 7200, 6600, 7920, 6600, 7200], "num"),
    ("HEADCOUNT", "تعداد نیروی انسانی مستقر", "نفر",
     [12, 18, 14, 24, 16, 22], "num"),
    ("POWER", "توان نصب‌شده", "کیلووات",
     [350, 900, 650, 1800, 900, 700], "num"),
    ("AREA", "مساحت اشغالی", "m2",
     [1800, 1400, 1200, 2200, 1300, 2600], "num"),
    ("WATER", "مصرف آب در سال (طراحی)", "m3/سال",
     [3000, 24000, 18000, 6000, 4000, 9000], "num"),
]

TECH_RAMP = [0, 0.6, 0.8, 0.9, 1, 1, 1, 1, 1, 1, 1, 1, 1]

# ------------------------------------------------------------------ IN_SALES
PRODUCTS = [
    "شمش سرب خالص (Pb 99.97)",
    "شمش آلیاژ سرب (Pb-Sb / Pb-Ca)",
    "گرانول پلی‌پروپیلن بازیافتی",
    "سولفات سدیم (محصول جانبی)",
]

SALES_ROWS = [
    ("SAL.price", "قیمت پایه (سال پایه)", "میلیون ریال/تن", [1050, 1082, 325, 60], "num"),
    ("SAL.lme", "پیوند قیمت به LME (1=بله/0=خیر)", "1/0", [1, 1, 0, 0], "idx"),
    ("SAL.lmef", "ضریب تعدیل نسبت به LME", "ضریب", [1, 1.03, 1, 1], "num2"),
    ("SAL.esc", "نرخ رشد سالانه قیمت داخلی", "درصد", [0.20, 0.20, 0.22, 0.20], "pct"),
    ("SAL.fxsh", "سهم ارزی (صادراتی/شاخص ارز)", "درصد", [0.35, 0.35, 0.10, 0.50], "pct"),
    ("SAL.disc", "نرخ تخفیف و برگشت از فروش", "درصد", [0.015, 0.015, 0.03, 0.05], "pct2"),
    ("SAL.real", "ضریب تحقق فروش (فروش/تولید)", "ضریب", [1, 1, 0.98, 0.95], "num2"),
]

SALES_SCALARS = [
    ("SAL.disc_late", "نرخ برگشت از فروش و ادعاهای کیفی", "درصد", 0.01, "pct2"),
    ("SAL.price_rev", "دوره بازنگری قیمت‌ها (ماه)", "ماه", 3, "num"),
    ("SAL.contract", "سهم فروش قراردادی بلندمدت", "درصد", 0.40, "pct"),
    ("SAL.spot", "سهم فروش نقدی/فوری", "درصد", 0.60, "pct"),
]

# ------------------------------------------------------------------- IN_COST
COST_SCALARS = [
    ("CST.fx_var", "سهم ارزی هزینه‌های متغیر", "درصد", 0.20, "pct"),
    ("CST.fx_fix", "سهم ارزی هزینه‌های ثابت", "درصد", 0.05, "pct"),
    ("CST.months", "ماه‌های پرداخت حقوق در سال", "ماه", 12, "num"),
    ("CST.bonus", "ضریب مزایا، عیدی و سنوات", "ضریب", 1.08, "num2"),
]

PAYROLL = [
    ("PAY1", "تولید مستقیم", 96, 250),
    ("PAY2", "تولید غیرمستقیم", 24, 320),
    ("PAY3", "نگهداری و تعمیرات", 18, 300),
    ("PAY4", "اداری و مالی", 22, 380),
    ("PAY5", "فروش و بازاریابی", 8, 350),
]

SERVICE_CENTERS = [
    ("SC1", "نگهداری و تعمیرات مرکزی", "MACHINE_HOURS", 300000, "F"),
    ("SC2", "منابع انسانی و خدمات پشتیبانی", "HEADCOUNT", 100000, "F"),
    ("SC3", "لجستیک و انبارداری", "THROUGHPUT", 150000, "F"),
    ("SC4", "تصفیه‌خانه و محیط‌زیست", "WATER", 250000, "V"),
    ("SC5", "آزمایشگاه و کنترل کیفیت", "THROUGHPUT", 80000, "V"),
]

# (کد, شرح, طبقه, مبنا, پارامتر, شاخص تعدیل, مبنای تخصیص, رفتار)
COST_ITEMS = [
    ("DPC-01", "باتری ضایعاتی (مواد اولیه اصلی)", "میلیون ریال/تن", "DPC", "SCRAP_LME", 348, "NONE", "THROUGHPUT", "V"),
    ("DPC-02", "مواد مصرفی و واکنشگرها", "میلیون ریال/تن", "DPC", "PER_TON_BATT", 15, "CPI", "THROUGHPUT", "V"),
    ("DPC-03", "دستمزد مستقیم تولید", "─", "DPC", "PAY", "PAY:1", "WAGE", "HEADCOUNT", "V"),
    ("DPC-04", "مواد بسته‌بندی", "میلیون ریال/تن", "DPC", "PER_TON_PROD", 4, "CPI", "THROUGHPUT", "V"),
    ("DPC-05", "حمل‌ونقل ورودی مواد", "میلیون ریال/تن", "DPC", "PER_TON_BATT", 6, "CPI", "THROUGHPUT", "V"),
    ("DPC-06", "انرژی الکتریکی فرآیندی", "میلیون ریال/kWh", "DPC", "PER_KWH", 0.006, "ENERGY", "MACHINE_HOURS", "V"),
    ("OPC-01", "سوخت و مواد احیاکننده", "میلیون ریال/تن", "OPC", "PER_TON_BATT", 12.5, "ENERGY", "THROUGHPUT", "V"),
    ("OPC-02", "آب و پساب", "میلیون ریال/m3", "OPC", "PER_M3", 0.25, "CPI", "WATER", "V"),
    ("OPC-03", "دفع پسماند و لجن", "میلیون ریال/تن", "OPC", "PER_TON_BATT", 7.5, "CPI", "THROUGHPUT", "V"),
    ("OPC-04", "نگهداری و تعمیرات متغیر", "میلیون ریال/تن", "OPC", "PER_TON_BATT", 9, "CPI", "MACHINE_HOURS", "V"),
    ("OPC-05", "مرکز خدماتی: نگهداری و تعمیرات", "─", "OPC", "SRV", "SRV:1", "CPI", "SRV", "F"),
    ("OPC-06", "مرکز خدماتی: لجستیک و انبارداری", "─", "OPC", "SRV", "SRV:3", "CPI", "SRV", "F"),
    ("OPC-07", "مرکز خدماتی: تصفیه‌خانه", "─", "OPC", "SRV", "SRV:4", "CPI", "SRV", "V"),
    ("OPC-08", "مرکز خدماتی: آزمایشگاه و QC", "─", "OPC", "SRV", "SRV:5", "CPI", "SRV", "V"),
    ("OPC-09", "هزینه فروش، توزیع و کارمزد فروش", "درصد از درآمد", "OPC", "PCT_REV", 0.012, "NONE", "THROUGHPUT", "V"),
    ("FPC-01", "استهلاک دارایی‌های ثابت", "─", "FPC", "DEP", "DEP", "NONE", "POWER", "F"),
    ("FPC-02", "بیمه دارایی‌ها", "درصد ارزش دفتری", "FPC", "PCT_ASSET", 0.006, "NONE", "POWER", "F"),
    ("FPC-03", "اجاره و حق‌الامتیاز", "میلیون ریال/سال", "FPC", "ANNUAL", 80000, "CPI", "AREA", "F"),
    ("FPC-04", "حقوق غیرمستقیم تولید", "─", "FPC", "PAY", "PAY:2", "WAGE", "HEADCOUNT", "F"),
    ("FPC-05", "حقوق نگهداری و تعمیرات", "─", "FPC", "PAY", "PAY:3", "WAGE", "HEADCOUNT", "F"),
    ("FPC-06", "حقوق اداری و مالی", "─", "FPC", "PAY", "PAY:4", "WAGE", "HEADCOUNT", "F"),
    ("FPC-07", "حقوق فروش و بازاریابی", "─", "FPC", "PAY", "PAY:5", "WAGE", "HEADCOUNT", "F"),
    ("FPC-08", "هزینه‌های عمومی و اداری غیرپرسنلی", "میلیون ریال/سال", "FPC", "ANNUAL", 300000, "CPI", "HEADCOUNT", "F"),
    ("FPC-09", "تبلیغات و بازاریابی ثابت", "میلیون ریال/سال", "FPC", "ANNUAL", 120000, "CPI", "AREA", "F"),
    ("FPC-10", "بار پایه انرژی", "میلیون ریال/سال", "FPC", "ANNUAL", 400000, "ENERGY", "POWER", "F"),
    ("FPC-11", "مرکز خدماتی: منابع انسانی", "─", "FPC", "SRV", "SRV:2", "CPI", "SRV", "F"),
    ("FPC-12", "سایر هزینه‌های ثابت", "میلیون ریال/سال", "FPC", "ANNUAL", 150000, "CPI", "AREA", "F"),
]

COST_NOTES = [
    "راهنما: مبنای محاسبه ← SCRAP_LME (خرید مواد اولیه، پیوند به LME×ارز)، PER_TON_BATT، PER_TON_PROD،",
    "PER_KWH، PER_M3، ANNUAL (مبلغ سالانه)، PCT_REV (درصد از درآمد)، PCT_ASSET (درصد ارزش دفتری)،",
    "DEP (استهلاک از موتور دارایی‌ها)، PAY:n (طبقه حقوقی n)، SRV:n (مرکز خدماتی n).",
    "شاخص تعدیل: CPI / WAGE / ENERGY / PPI / FX / NONE. مبنای تخصیص: THROUGHPUT / MACHINE_HOURS /",
    "HEADCOUNT / POWER / AREA / WATER / SRV (تخصیص از جدول مراکز خدماتی).",
]

# --------------------------------------------------------------------- IN_WC
WC_SCALARS_A = [
    ("WC.dio_rm", "موجودی مواد اولیه (باتری ضایعاتی)", "روز", 25, "num"),
    ("WC.dio_wip", "موجودی کالای در جریان ساخت", "روز", 8, "num"),
    ("WC.dio_fg", "موجودی محصول نهایی", "روز", 15, "num"),
    ("WC.dio_cons", "موجودی مواد مصرفی و قطعات یدکی", "روز", 30, "num"),
    ("WC.dso", "دوره وصول مطالبات (DSO)", "روز", 45, "num"),
    ("WC.dpo", "دوره پرداخت بدهی‌ها (DPO)", "روز", 40, "num"),
    ("WC.dpo_tax", "دوره پرداخت مالیات", "روز", 30, "num"),
    ("WC.cashdays", "ذخیره نقدی عملیاتی (روز هزینه نقدی)", "روز", 15, "num"),
]

WC_SCALARS_B = [
    ("WC.wc_loan_share", "سهم سرمایه در گردش تأمین‌شده از تسهیلات کوتاه‌مدت", "درصد", 0.30, "pct"),
    ("WC.bad_debt", "نرخ ذخیره مطالبات مشکوک‌الوصول", "درصد", 0.01, "pct2"),
    ("WC.inv_obs", "ضریب کاهش ارزش موجودی (کهنگی)", "درصد", 0.005, "pct2"),
]

# -------------------------------------------------------------------- IN_FIN
# نام، مبلغ کل، سهم ارزی، عمر مفید، دوره جایگزینی، درصد جایگزینی، ارزش اسقاط، تزریق t=0..t=2
CAPEX_ITEMS = [
    ("CAP-01", "زمین", 300000, 0.00, 0, 0, 0.00, 1.00, [1.00, 0.00, 0.00]),
    ("CAP-02", "محوطه‌سازی و ساختمان", 1600000, 0.10, 30, 0, 0.00, 0.30, [0.75, 0.25, 0.00]),
    ("CAP-03", "ماشین‌آلات وارداتی", 4800000, 0.90, 15, 0, 0.60, 0.10, [0.75, 0.25, 0.00]),
    ("CAP-04", "ماشین‌آلات داخلی", 1900000, 0.25, 12, 12, 0.70, 0.10, [0.75, 0.25, 0.00]),
    ("CAP-05", "تأسیسات و تجهیزات جانبی", 1200000, 0.40, 15, 0, 0.50, 0.10, [0.75, 0.25, 0.00]),
    ("CAP-06", "تصفیه‌خانه و تجهیزات زیست‌محیطی", 1050000, 0.30, 12, 12, 0.70, 0.10, [0.75, 0.25, 0.00]),
    ("CAP-07", "تجهیزات حمل‌ونقل داخلی", 700000, 0.20, 8, 8, 0.80, 0.15, [0.60, 0.40, 0.00]),
    ("CAP-08", "آزمایشگاه و ابزار دقیق", 450000, 0.60, 10, 10, 0.60, 0.10, [0.50, 0.50, 0.00]),
    ("CAP-09", "مهندسی، نصب و راه‌اندازی", 900000, 0.10, 15, 0, 0.00, 0.00, [0.60, 0.40, 0.00]),
    ("CAP-10", "هزینه‌های پیش از بهره‌برداری", 420000, 0.05, 5, 5, 0.50, 0.00, [0.30, 0.50, 0.20]),
    ("CAP-11", "ذخیره احتیاطی فیزیکی", 1200000, 0.40, 15, 0, 0.00, 0.10, [0.40, 0.40, 0.20]),
]

FIN_SCALARS = [
    ("FIN.debt_share", "سهم تسهیلات بلندمدت از CAPEX", "درصد", 0.60, "pct"),
    ("FIN.kd", "نرخ سود تسهیلات بلندمدت", "درصد", 0.22, "pct"),
    ("FIN.fee_arr", "کارمزد تخصیص/تعهد یکجا", "درصد", 0.015, "pct2"),
    ("FIN.fee_com", "کارمزد تعهد روی مانده استفاده‌نشده", "درصد", 0.005, "pct2"),
    ("FIN.repay_start", "شاخص سال شروع بازپرداخت (t)", "شاخص", 2, "idx"),
    ("FIN.repay_years", "دوره بازپرداخت (سال)", "سال", 8, "num"),
    ("FIN.repay_method", "روش بازپرداخت (1=اقساط مساوی اصل / 2=آنویتی)", "1/2", 1, "idx"),
    ("FIN.kd_st", "نرخ سود تسهیلات کوتاه‌مدت (سرمایه در گردش)", "درصد", 0.24, "pct"),
    ("FIN.min_dscr", "حداقل DSCR قراردادی", "ضریب", 1.30, "num2"),
    ("FIN.payout", "سیاست تقسیم سود (درصد از سود خالص)", "درصد", 0.40, "pct"),
    ("FIN.i_income", "نرخ سود سپرده/درآمد غیرعملیاتی نقد", "درصد", 0.10, "pct"),
]

# ------------------------------------------------------------------- IN_SCEN
SCEN_ROWS = [
    ("VOL", "ضریب حجم تولید و فروش", [1.00, 1.08, 0.88], "num2"),
    ("PRICE", "ضریب قیمت فروش محصولات", [1.00, 1.10, 0.90], "num2"),
    ("RM", "ضریب قیمت خرید باتری ضایعاتی", [1.00, 0.95, 1.08], "num2"),
    ("VAR", "ضریب سایر هزینه‌های متغیر", [1.00, 0.97, 1.10], "num2"),
    ("FIX", "ضریب هزینه‌های ثابت", [1.00, 0.95, 1.12], "num2"),
    ("CAPEX", "ضریب سرمایه‌گذاری ثابت", [1.00, 0.95, 1.15], "num2"),
    ("FX", "ضریب نرخ ارز", [1.00, 1.05, 0.95], "num2"),
    ("WACC", "تعدیل نرخ تنزیل", [0.00, -0.02, 0.03], "pct"),
    ("KD", "تعدیل نرخ سود تسهیلات", [0.00, -0.01, 0.03], "pct"),
]

SCEN_AXES = [
    ("SCEN.ax_rate", "انحراف نرخ بهره/تنزیل", [-0.03, -0.015, 0, 0.015, 0.03], "pct"),
    ("SCEN.ax_fx", "انحراف نرخ ارز", [-0.20, -0.10, 0, 0.10, 0.20], "pct"),
    ("SCEN.ax_price", "انحراف قیمت فروش", [-0.15, -0.075, 0, 0.075, 0.15], "pct"),
]

SCEN_NOTES = [
    "پایه: مفروضاتِ مندرج در شیت‌های ورودی بدون تعدیل (تمام ضرایب برابر یک یا صفر).",
    "خوش‌بینانه: رشد حجم و قیمت، کنترل هزینه‌ها و دسترسی بهتر به تأمین مالی.",
    "بدبینانه: افت حجم و قیمت، افزایش هزینه‌ها، افزایش CAPEX و نرخ تأمین مالی.",
]


def build(b):
    """ساختِ هفت شیتِ ورودی"""
    _macro(b)
    _tech(b)
    _sales(b)
    _cost(b)
    _wc(b)
    _fin(b)
    _scen(b)


# ---------------------------------------------------------------------------
def _macro(b):
    sh = b.sheet("IN_MACRO", "۱. فرضیات کلان اقتصادی",
                 "تمام مفروضات اقتصاد کلان در این شیت وارد می‌شود؛ هیچ فرمولی مجاز نیست. "
                 "واحد پول مدل: میلیون ریال.", "input")
    sh.section("الف) پارامترهای کنترلی و ثابت مدل")
    for key, lab, unit, val, fmt in MACRO_SCALARS:
        sh.add_scalar(key, lab, unit, val, fmt)
    sh.gap()
    sh.section("ب) مسیرهای سالانهٔ اقتصاد کلان (ورودی عددی)")
    for key, lab, unit, vals, fmt in MACRO_SERIES:
        sh.add(key, lab, unit, vals, fmt)
    sh.gap()
    sh.section("ج) یادداشت‌های روش‌شناختی")
    for n in MACRO_NOTES:
        sh.note(n)


def _tech(b):
    sh = b.sheet("IN_TECH", "۲. پارامترهای فنی خط تولید (ایستگاه‌های ۱ تا ۶)",
                 "ظرفیت اسمی و مؤثر، راندمان، ضایعات، زمان‌های توقف و محرک‌های تخصیص هزینه.",
                 "input", time_axis=False)
    sh.r = 7
    sh.section("الف) پارامترهای سطح کارخانه")
    for key, lab, unit, val, fmt in TECH_SCALARS_A:
        sh.add_scalar(key, lab, unit, val, fmt)
    sh.gap()
    sh.section("ب) عوامل بازیابی و تبدیل مواد")
    for key, lab, unit, val, fmt in TECH_SCALARS_B:
        sh.add_scalar(key, lab, unit, val, fmt)
    sh.gap()
    sh.section("ج) ترکیب جرمی باتری ورودی (سهم از هر تن)")
    for key, lab, unit, val, fmt in TECH_SCALARS_C:
        sh.add_scalar(key, lab, unit, val, fmt)
    sh.gap()
    sh.section("د) جدول پارامترهای ایستگاه‌ها")
    sh.ws.cell(sh.r, 1, "ایستگاه")
    sh.header(["S%d" % (i + 1) for i in range(spec.NS)])
    for i, name in enumerate(STATIONS):
        sh.ws.cell(sh.r, spec.C0 + i, name)
    sh.r += 1
    for key, lab, unit, vals, fmt in STATION_ROWS:
        sh.table(key, lab, unit, vals, fmt, n=spec.NS)
    sh.gap()
    sh.section("ه‍) جدول محرک‌های تخصیص هزینه به ایستگاه‌ها")
    sh.ws.cell(sh.r, 1, "محرک / ایستگاه")
    sh.header(["S%d" % (i + 1) for i in range(spec.NS)])
    for i, name in enumerate(STATIONS):
        sh.ws.cell(sh.r, spec.C0 + i, name)
    sh.r += 1
    for key, lab, unit, vals, fmt in DRIVER_ROWS:
        sh.table(key, lab, unit, vals, fmt, n=spec.NS)
    sh.gap()
    sh.note("توضیح: این جدول مبنای توزیع هزینه‌ها بین ایستگاه‌های ۱ تا ۶ است.")
    sh.gap()
    sh.section("و) منحنی بهره‌برداری از ظرفیت (سهم از ظرفیت اسمی)")
    sh.local_axis()
    sh.add("TECH.ramp", "ضریب دستیابی به ظرفیت اسمی", "درصد", TECH_RAMP, "pct")


def _sales(b):
    sh = b.sheet("IN_SALES", "۳. مفروضات قیمت‌گذاری و فروش",
                 "قیمت پایه، پیوند به قیمت جهانی، ترکیب سبد محصولات و سیاست‌های تخفیف و اعتبار.",
                 "input", time_axis=False)
    sh.r = 7
    sh.section("الف) ویژگی‌های سبد محصولات")
    sh.ws.cell(8, 1, "محصول")
    sh.header(["P%d" % (i + 1) for i in range(spec.NP)], row=8)
    for i, name in enumerate(PRODUCTS):
        sh.ws.cell(9, spec.C0 + i, name)
    sh.r = 10
    for key, lab, unit, vals, fmt in SALES_ROWS:
        sh.table(key, lab, unit, vals, fmt, n=spec.NP)
    sh.gap()
    sh.section("ب) سیاست‌های اعتباری و فروش")
    for key, lab, unit, val, fmt in SALES_SCALARS:
        sh.add_scalar(key, lab, unit, val, fmt)
    sh.gap()
    sh.note("توضیح: «سهم ارزی» نشان می‌دهد چه بخشی از درآمدِ هر محصول با نرخ ارز تعدیل می‌شود "
            "(مبنای تحلیل حساسیت ارزی).")


def _cost(b):
    sh = b.sheet("IN_COST", "۴. پارامترهای هزینه‌ای (FPC / OPC / DPC)",
                 "نرخ‌های واحد، پرسنل، مراکز هزینهٔ خدماتی و ماتریس اقلام هزینه با مبنای محاسبه و تخصیص.",
                 "input", time_axis=False)
    sh.r = 7
    sh.section("الف) نرخ‌ها و پارامترهای پایه")
    for key, lab, unit, val, fmt in COST_SCALARS:
        sh.add_scalar(key, lab, unit, val, fmt)
    sh.gap()
    sh.section("ب) پرسنل (تعداد و حقوق ماهانهٔ پایه)")
    sh.header(["کد طبقه", "تعداد نفر", "حقوق ماهانه", "واحد"], row=14)
    sh.r = 15
    for code, lab, head, salary in PAYROLL:
        sh.add_scalar(code, lab, None, head, "num")
        c = sh.ws.cell(sh.r - 1, 5, salary)
        c.number_format = common.FMT["num"]
        c.fill = common.INPUT_FILL
        c.border = common.BOX
        sh.ws.cell(sh.r - 1, 6, "نفر / میلیون ریال").font = common.UNIT_FONT
    sh.gap()
    sh.section("ج) مراکز هزینهٔ خدماتی")
    sh.header(["نام مرکز", "محرک تخصیص", "هزینه سالانهٔ پایه", "رفتار", "ترتیب"], row=22)
    sh.r = 23
    for i, (code, name, driver, amount, behav) in enumerate(SERVICE_CENTERS):
        sh.ws.cell(sh.r, 1, name)
        c = sh.ws.cell(sh.r, 2, code)
        c.font = common.KEY_FONT
        b.keys[(sh.name, code)] = sh.r
        for j, v in enumerate([name, driver, amount, behav, i + 1]):
            cc = sh.ws.cell(sh.r, 4 + j, v)
            cc.fill = common.INPUT_FILL
            cc.border = common.BOX
            if j == 2:
                cc.number_format = common.FMT["num"]
        sh.r += 1
    sh.gap()
    sh.note("رفتار هزینه: V = متغیر (تابع حجم)، F = ثابت (سالانه).")
    sh.gap()
    sh.section("د) ماتریس اقلام هزینه (مبنای محاسبه، تعدیل و تخصیص)")
    sh.header(["طبقه", "مبنای محاسبه", "پارامتر / مرجع", "شاخص تعدیل", "مبنای تخصیص", "رفتار"], row=32)
    sh.r = 33
    first = sh.r
    for code, lab, unit, cls, basis, param, index, alloc, behav in COST_ITEMS:
        sh.ws.cell(sh.r, 1, lab)
        c = sh.ws.cell(sh.r, 2, code)
        c.font = common.KEY_FONT
        b.keys[(sh.name, code)] = sh.r
        sh.ws.cell(sh.r, 3, unit).font = common.UNIT_FONT
        for j, v in enumerate([cls, basis, param, index, alloc, behav]):
            cc = sh.ws.cell(sh.r, 4 + j, v)
            cc.fill = common.INPUT_FILL
            cc.border = common.BOX
            if j == 2 and isinstance(v, (int, float)):
                cc.number_format = common.FMT["num3"] if v < 1 else common.FMT["num"]
        sh.r += 1
    last = sh.r - 1
    sh.gap()
    for n in COST_NOTES:
        sh.note(n)
    # اعتبارسنجی‌های داده برای ستون‌های ماتریس هزینه
    for col, items in ((4, "DPC,OPC,FPC"),
                       (5, "SCRAP_LME,PER_TON_BATT,PER_TON_PROD,PER_KWH,PER_M3,ANNUAL,"
                           "PCT_REV,PCT_ASSET,DEP,PAY,SRV"),
                       (7, "CPI,WAGE,ENERGY,PPI,FX,NONE"),
                       (8, "THROUGHPUT,MACHINE_HOURS,HEADCOUNT,POWER,AREA,WATER,SRV"),
                       (9, "V,F")):
        dv = DataValidation(type="list", formula1='"%s"' % items, allow_blank=True)
        dv.error = "مقدار باید یکی از گزینه‌های مجاز باشد"
        dv.errorTitle = "ورودی نامعتبر"
        sh.ws.add_data_validation(dv)
        dv.add("%s%d:%s%d" % (chr(64 + col), first, chr(64 + col), last))


def _wc(b):
    sh = b.sheet("IN_WC", "۵. سیاست‌های سرمایه در گردش و چرخه نقدینگی",
                 "روزهای موجودی، دریافتنی و پرداختنی برای محاسبهٔ چرخهٔ تبدیل نقد (CCC).",
                 "input", time_axis=False)
    sh.r = 7
    sh.section("الف) دوره‌های سرمایه در گردش (روز)")
    for key, lab, unit, val, fmt in WC_SCALARS_A:
        sh.add_scalar(key, lab, unit, val, fmt)
    sh.gap()
    sh.section("ب) سیاست‌های تأمین مالی سرمایه در گردش")
    for key, lab, unit, val, fmt in WC_SCALARS_B:
        sh.add_scalar(key, lab, unit, val, fmt)


def _fin(b):
    sh = b.sheet("IN_FIN", "۶. ساختار سرمایه، تأمین مالی و برنامه CAPEX",
                 "درصد بدهی و حقوق صاحبان سهام، نرخ‌های بهره، دوره بازپرداخت، کارمزدها و برنامه تزریق سرمایه.",
                 "input", time_axis=False)
    sh.r = 7
    sh.section("الف) برنامه سرمایه‌گذاری ثابت (CAPEX)")
    sh.header(["مبلغ کل", "سهم ارزی", "عمر مفید", "دوره جایگزینی", "درصد جایگزینی",
               "ارزش اسقاط", "تزریق t=0", "تزریق t=1", "تزریق t=2"], row=8)
    sh.r = 9
    for code, lab, total, fxsh, life, repl, repl_pct, salvage, inject in CAPEX_ITEMS:
        sh.ws.cell(sh.r, 1, lab)
        c = sh.ws.cell(sh.r, 2, code)
        c.font = common.KEY_FONT
        b.keys[(sh.name, code)] = sh.r
        sh.ws.cell(sh.r, 3, "میلیون ریال").font = common.UNIT_FONT
        vals = [total, fxsh, life, repl, repl_pct, salvage] + list(inject)
        fmts = ["num", "pct", "num", "num", "pct", "num2", "pct", "pct", "pct"]
        for j, v in enumerate(vals):
            cc = sh.ws.cell(sh.r, 4 + j, v)
            cc.number_format = common.FMT[fmts[j]]
            cc.fill = common.INPUT_FILL
            cc.border = common.BOX
        sh.r += 1
    sh.gap()
    sh.note("توضیح: عمر مفید صفر = مستهلک نمی‌شود (زمین). دوره جایگزینی صفر = جایگزینی در افق مدل ندارد.")
    sh.gap()
    sh.section("ب) ساختار تأمین مالی")
    for key, lab, unit, val, fmt in FIN_SCALARS:
        sh.add_scalar(key, lab, unit, val, fmt)


def _scen(b):
    sh = b.sheet("IN_SCEN", "۷. تعریف سناریوها و محورهای حساسیت",
                 "ضرایب تعدیل مفروضات کلیدی در سناریوهای پایه / خوش‌بینانه / بدبینانه و محورهای ماتریس حساسیت.",
                 "input", time_axis=False)
    sh.r = 7
    sh.section("الف) انتخاب سناریوی فعال")
    sh.add_scalar("SCEN.active_in", "سناریوی فعال (۱=پایه، ۲=خوش‌بینانه، ۳=بدبینانه)", "شاخص", 1, "idx")
    sh.note("تغییر این سلول کل مدل را به‌صورت زنجیره‌ای به‌روزرسانی می‌کند (ورودی ← موتور ← خروجی).")
    sh.gap()
    sh.section("ب) جدول ضرایب سناریوها")
    sh.header(["پایه", "خوش‌بینانه", "بدبینانه"], row=12)
    sh.r = 13
    for key, lab, vals, fmt in SCEN_ROWS:
        sh.table(key, lab, "ضریب" if fmt == "num2" else "واحد درصد", vals, fmt, n=3)
    sh.gap()
    sh.section("ج) محورهای ماتریس حساسیت")
    for key, lab, vals, fmt in SCEN_AXES:
        sh.table(key, lab, "درصد", vals, fmt, n=5)
    sh.gap()
    sh.section("د) شرح سناریوها")
    for n in SCEN_NOTES:
        sh.note(n)
